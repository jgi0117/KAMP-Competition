"""Append the experiment design and clearly labelled illustrative plots to the report.

This script enumerates proposed trials and creates synthetic visual examples.
It does not train forecasting models or produce measured performance results.
"""
from __future__ import annotations

import csv
import itertools
import json
import re
import shutil
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / 'report_assets'
ASSETS.mkdir(exist_ok=True)
MARKER = '## 5. 하이퍼파라미터 탐색 및 층별 미세조정 실험 설계'
TITLE = '05_전력사용량_모델후보_평가지표_선정'
NOTE = '가상 데이터 · 실제 학습 결과 또는 성능 예측값이 아님'
plt.rcParams.update({
    'font.family': 'Malgun Gothic', 'axes.unicode_minus': False,
    'font.size': 10, 'axes.spines.top': False, 'axes.spines.right': False,
    'axes.titleweight': 'bold', 'figure.facecolor': 'white',
    'savefig.facecolor': 'white', 'svg.fonttype': 'path',
})
COLORS = ['#2563eb', '#0f9d8a', '#d88624', '#8b5cf6', '#db647c', '#475569']

GRIDS = {
    'XGBoost': {'max_depth': [3, 5, 7], 'learning_rate': [.01, .05, .1],
                'n_estimators': [200, 500, 1000], 'min_child_weight': [1, 5],
                'colsample_bytree': [.8, 1.]},
    'LightGBM': {'num_leaves': [7, 15, 31], 'learning_rate': [.01, .05, .1],
                 'n_estimators': [200, 500, 1000], 'min_child_samples': [20, 50],
                 'colsample_bytree': [.8, 1.]},
    'LSTM': {'hidden_size': [32, 64, 128], 'num_layers': [1, 2],
             'learning_rate': [.0001, .0003, .001], 'head_dropout': [0., .2]},
}
SCOPES = {
    'Chronos-2': {'total_blocks': 12, 'last_k': [1, 2, 3, 6, 9, 12],
                 'blocks': 'encoder.block',
                 'head': ['output_patch_embedding', 'encoder.final_layer_norm'],
                 'status': 'official full tuning; selective freezing needs a custom training path'},
    'TimesFM 3.0': {'total_blocks': 20, 'last_k': [1, 2, 5, 10, 15, 20],
                   'blocks': 'transformer_stack.layers', 'head': ['output_head'],
                   'status': 'inference-only public implementation; custom training and gradient validation required'},
    'Moirai 2.0': {'total_blocks': 6, 'last_k': [1, 2, 3, 4, 5, 6],
                  'blocks': 'encoder.layers', 'head': ['out_proj', 'encoder.norm'],
                  'status': 'training_mode forward exists; Moirai2-specific training adapter required'},
}


def write_csv(name, rows):
    rows = list(rows)
    with (ASSETS / name).open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def enumerate_plan():
    trials = []
    for model, grid in GRIDS.items():
        for i, values in enumerate(itertools.product(*grid.values()), 1):
            trials.append({'model': model, 'trial_id': f'{model}-{i:03}',
                           'parameters': json.dumps(dict(zip(grid, values))), 'status': 'planned'})
    write_csv('planned_grid_trials.csv', trials)
    layers = []
    for model, cfg in SCOPES.items():
        settings = [('ZS', None), ('F0', 0)] + [(f'F{i}', k) for i, k in enumerate(cfg['last_k'], 1)] + [('F7', 'all')]
        for setting, k in settings:
            layers.append({'model': model, 'setting': setting, 'last_k_blocks': k,
                           'total_blocks': cfg['total_blocks'], 'implementation_status': cfg['status'],
                           'status': 'planned'})
    write_csv('planned_layer_trials.csv', layers)
    plan = {'status': 'proposal_not_trained', 'checked_on': '2026-10-05',
            'grid': GRIDS, 'grid_counts': {m: int(np.prod([len(x) for x in g.values()])) for m, g in GRIDS.items()},
            'layer_scopes': SCOPES, 'fine_tune_learning_rates': [1e-6, 1e-5, 1e-4],
            'protocol': {'proposed_context_hours': 168, 'proposed_horizon_hours': 24,
                         'proposed_test_fraction': .2, 'expanding_validation_folds': 3,
                         'confirmation_seeds': [17, 42, 73], 'peak_train_quantile': .95,
                         'point_prediction': 'median for quantile models',
                         'selection': 'validation RMSE within 2% of best, then peak MAE, then MAE and time'},
            'figures': {'data_type': 'synthetic', 'purpose': 'preview of evaluation outputs, not expected model rankings'}}
    (ASSETS / 'experiment_plan.json').write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding='utf-8')


def finish(fig, stem):
    fig.text(.5, .012, NOTE, ha='center', va='bottom', color='#b45309', fontsize=10, fontweight='bold')
    fig.savefig(ASSETS / (stem + '.png'), dpi=180)
    fig.savefig(ASSETS / (stem + '.svg'))
    plt.close(fig)


def make_figures():
    # Every numeric series below is hand-constructed or simulated for illustration.
    matrices = [np.array([[18.8, 16.1, 17.4], [17.2, 14.5, 15.9], [17.6, 15.2, 17.8]]),
                np.array([[19., 16.4, 17.], [17.6, 14.8, 15.5], [17.4, 15.3, 17.2]]),
                np.array([[20.1, 18.2, 18.9], [18.7, 16.2, 16.8], [18.4, 16.7, 18.5]])]
    fig, axes = plt.subplots(1, 3, figsize=(13.8, 5.2))
    rows = []
    for ax, model, mat, ys, xs, ylabel in zip(axes, GRIDS, matrices,
            [[3, 5, 7], [7, 15, 31], [32, 64, 128]],
            [['.01', '.05', '.10'], ['.01', '.05', '.10'], ['.0001', '.0003', '.001']],
            ['max_depth', 'num_leaves', 'hidden_size']):
        im = ax.imshow(mat, cmap='YlGnBu', vmin=14, vmax=21, aspect='auto')
        ax.set(xticks=range(3), xticklabels=xs, yticks=range(3), yticklabels=ys,
               xlabel='learning_rate', ylabel=ylabel, title=model)
        for i, j in itertools.product(range(3), repeat=2):
            ax.text(j, i, f'{mat[i,j]:.1f}', ha='center', va='center', color='white' if mat[i,j] > 18 else '#172554')
            rows.append({'data_type': 'synthetic', 'model': model, ylabel: ys[i],
                         'learning_rate': xs[j], 'validation_RMSE_example': mat[i,j]})
        i, j = np.unravel_index(mat.argmin(), mat.shape)
        ax.add_patch(plt.Rectangle((j-.48, i-.48), .96, .96, fill=False, ec='#ef4444', lw=2.5))
    fig.suptitle('그림 1. Grid Search 결과표 예시 — 검증 RMSE가 낮은 조합 찾기', y=.98, fontsize=15, fontweight='bold')
    fig.text(.5, .09, '고정: 트리 500개 / colsample_bytree=1.0 / XGB min_child_weight=1 / LGB min_child_samples=20\nLSTM: 1층 / 출력부 dropout=0.0 · 나머지는 본문 고정 설정', ha='center', color='#475569', fontsize=9)
    fig.subplots_adjust(left=.06, right=.89, top=.81, bottom=.23, wspace=.42)
    cax = fig.add_axes([.92, .23, .016, .58])
    fig.colorbar(im, cax=cax, label='검증 RMSE (가상 전력 단위)')
    finish(fig, '01_grid_search_example')
    # Avoid variable CSV columns by expressing both grid coordinates generically.
    write_csv('illustrative_grid.csv', [{'data_type': r['data_type'], 'model': r['model'],
        'structure_parameter': next(k for k in r if k in ['max_depth','num_leaves','hidden_size']),
        'structure_value': next(v for k,v in r.items() if k in ['max_depth','num_leaves','hidden_size']),
        'learning_rate': r['learning_rate'], 'validation_RMSE_example': r['validation_RMSE_example']} for r in rows])

    layer_vals = {
        'Chronos-2': [18.5, 17.1, 16.2, 15.5, 15.0, 14.6, 14.8, 15.3, 16.0],
        'TimesFM 3.0': [17.8, 17.1, 16.8, 16.4, 16.0, 15.7, 15.4, 15.2, 15.0],
        'Moirai 2.0': [19.4, 18.3, 17.6, 17.2, 17.1, 17.0, 17.1, 17.0, 17.2],
    }
    fig, axes = plt.subplots(1, 3, figsize=(13.8, 5.5))
    rows = []
    for ax, (model, vals), color in zip(axes, layer_vals.items(), COLORS):
        v = np.array(vals)
        error = np.array([0, .5, .4, .35, .3, .35, .4, .55, .65])
        train = np.r_[np.nan, np.linspace(14.8, 8.5, 8)]
        ax.axhline(v[0], c='#94a3b8', ls='--', lw=1, label='Zero-shot 검증 RMSE')
        ax.errorbar(range(9), v, yerr=error, color=color, marker='o', capsize=3, label='검증 RMSE')
        ax.plot(range(9), train, color='#64748b', ls=':', label='훈련 RMSE')
        ax.set(xticks=range(9), xticklabels=['ZS']+[f'F{i}' for i in range(8)], ylim=(7.5, 21),
               ylabel='RMSE (가상 전력 단위)', xlabel='학습 범위 (본문의 모델별 설정표)',
               title=model + (' *' if model != 'Chronos-2' else ''))
        ax.grid(axis='y', alpha=.2)
        ax.legend(fontsize=8, loc='lower left')
        for i in range(9):
            rows.append({'data_type':'synthetic', 'model': model, 'scope': 'ZS' if i==0 else f'F{i-1}',
                         'validation_RMSE':v[i], 'illustrative_sd': error[i],
                         'train_RMSE': '' if i==0 else train[i]})
    fig.suptitle('그림 2. 층별 Fine-tuning 추세 예시 — 과적합 / 지속 개선 / 정체', y=.98, fontsize=15, fontweight='bold')
    fig.text(.5, .09, '학습률 1e-5 고정 예시 · * 추가 학습 구현이 필요한 조건부 실험\n오차막대: 가상 반복 실험 표준편차', ha='center', color='#475569', fontsize=9)
    fig.subplots_adjust(left=.06, right=.98, top=.80, bottom=.23, wspace=.32)
    finish(fig, '02_layer_trends_example')
    write_csv('illustrative_layer_trends.csv', rows)

    names = ['XGBoost', 'LightGBM', 'LSTM', 'TimesFM 3.0 / ZS', 'TimesFM 3.0 / FT*',
             'Chronos-2 / ZS', 'Chronos-2 / FT', 'Moirai 2.0 / ZS', 'Moirai 2.0 / FT*']
    vals = np.array([[10.4,14.5,22.0],[10.2,14.8,19.0],[11.2,16.2,17.8],
                     [12.3,17.8,23.8],[10.7,15.0,20.8],[12.7,18.5,25.5],
                     [10.0,14.6,18.3],[13.5,19.4,26.3],[11.7,17.0,21.9]])
    times = np.array([2.4,1.4,13,0,20,0,9,0,4.5])
    fig, axes = plt.subplots(1, 3, figsize=(13.8, 6.4), sharey=True)
    for ax, i, color, title in zip(axes, range(3), COLORS, ['전체 MAE','전체 RMSE','피크 구간 MAE']):
        ax.barh(range(9), vals[:,i], color=[color if '/ ZS' not in n else '#cbd5e1' for n in names], height=.64)
        ax.set(yticks=range(9), yticklabels=names, xlabel='오차 (가상 전력 단위)', title=title, xlim=(0,29))
        for j, v in enumerate(vals[:,i]): ax.text(v+.4, j, f'{v:.1f}', va='center', fontsize=9)
        ax.grid(axis='x', alpha=.18)
    axes[0].invert_yaxis()
    fig.suptitle('그림 3. 후보별 성능 비교 예시 — 전체 오차와 피크 오차 함께 보기', y=.97, fontsize=15, fontweight='bold')
    fig.text(.5,.09,'각 FT는 검증 구간에서 선택한 설정 · * 구현 검증 후에만 결과표에 포함 · 후보 순위는 예시',ha='center',color='#475569')
    fig.subplots_adjust(left=.19, right=.96, top=.83, bottom=.20, wspace=.20)
    finish(fig, '03_model_comparison_example')
    write_csv('illustrative_model_comparison.csv', [{'data_type':'synthetic', 'model_variant':n,
        'MAE':v[0], 'RMSE':v[1], 'peak_MAE':v[2], 'fit_minutes_example': t}
        for n,v,t in zip(names, vals, times)])

    t = np.arange(72)
    target = 80 + 17*np.sin(2*np.pi*(t-6)/24) + 7*np.sin(2*np.pi*t/8)
    target += 57*np.exp(-((t-39)/2.2)**2) + 38*np.exp(-((t-62)/2.)**2)
    pa = target - 28*np.exp(-((t-39)/2.8)**2) - 20*np.exp(-((t-62)/2.8)**2) + 2*np.sin(t)
    pb = target - 6*np.exp(-((t-39)/2.8)**2) - 5*np.exp(-((t-62)/2.8)**2) + 6*np.sin(t/3)
    fig, axes = plt.subplots(2,1,figsize=(13.8,6.8),sharex=True,gridspec_kw={'height_ratios':[2,1]})
    axes[0].plot(t,target,c='#172554',lw=2,label='관측값 역할의 합성 시계열')
    axes[0].plot(t,pa,c=COLORS[0],lw=1.5,label='가상 후보 A')
    axes[0].plot(t,pb,c=COLORS[1],lw=1.5,label='가상 후보 B')
    axes[0].axhline(120,c='#b45309',ls='--',label='훈련 구간에서 정하는 피크 기준 (예시)')
    axes[0].fill_between(t,0,180,where=target>=120,color='#fbbf24',alpha=.17)
    axes[0].set(ylim=(45,170),ylabel='전력사용량 (가상 단위)')
    axes[0].legend(ncol=2,fontsize=9,loc='upper left')
    axes[1].plot(t,pa-target,c=COLORS[0],label='후보 A 오차')
    axes[1].plot(t,pb-target,c=COLORS[1],label='후보 B 오차')
    axes[1].axhline(0,c='#64748b',lw=.8)
    axes[1].set(ylabel='예측 - 관측',xlabel='동일 평가 구간의 경과 시간 (시간)')
    for ax in axes: ax.grid(alpha=.18)
    fig.suptitle('그림 4. 피크 예측 확대 예시 — 최고점 누락과 과소예측 확인',y=.98,fontsize=15,fontweight='bold')
    fig.subplots_adjust(left=.08,right=.97,top=.86,bottom=.13,hspace=.15)
    finish(fig,'04_peak_trace_example')
    write_csv('illustrative_peak_trace.csv',[{'data_type':'synthetic','hour':i,'synthetic_target':y,
        'candidate_A':a,'candidate_B':b} for i,y,a,b in zip(t,target,pa,pb)])

    fig, ax = plt.subplots(figsize=(11.8,6.5))
    for i,(n,v,tm) in enumerate(zip(names,vals,times)):
        ax.scatter(tm,v[1],s=65,color=COLORS[i%len(COLORS)],marker='s' if '/ ZS' in n else 'o')
        ax.annotate(n,(tm,v[1]),xytext=(8,6),textcoords='offset points',fontsize=9)
    ax.set(xlim=(-1,27),ylim=(13.5,21),xlabel='선택된 설정 1회의 추가 학습 시간 (분, 가상 값)',
           ylabel='검증 RMSE (가상 전력 단위)')
    ax.grid(alpha=.2)
    fig.suptitle('그림 5. 정확도와 학습 비용 비교 예시',y=.97,fontsize=15,fontweight='bold')
    fig.text(.5,.09,'Zero-shot의 추가 학습 시간은 0 · 모델 로딩 / 추론 / 전체 탐색 시간은 별도 기록',ha='center',color='#475569')
    fig.subplots_adjust(left=.10,right=.96,top=.85,bottom=.20)
    finish(fig,'05_cost_accuracy_example')


APPENDIX = r'''
## 5. 하이퍼파라미터 탐색 및 층별 미세조정 실험 설계

### 5.1 비교 조건과 선정 절차

아래 수치는 이 과제의 탐색 제안값이다. 논문에서 검증된 최적값이나 이번 데이터로 학습해 얻은 결과를 뜻하지 않는다. 파라미터의 의미와 구현 가능 여부는 공식 문서·모델 설정·소스 코드로 확인했다. 확인일은 2026년 10월 5일이다.

| 항목 | 제안하는 공통 조건 |
|---|---|
| 예측 과제 | 시간 단위 관측을 전제로 과거 168시간을 입력하고 이후 24시간을 예측한다. 실제 사용 시점에 맞춰 예측 길이를 바꾸면 모든 후보에 함께 적용한다. |
| 예측 시작 시점 | 동일한 시각에서 24시간 간격으로 예측해 평가 대상 구간이 중복되지 않게 한다. 트리 모델은 예측 시차별 직접 예측, LSTM은 24개 값을 출력하는 구조를 사용한다. |
| 공통 입력 | 먼저 전력 이력만으로 여섯 후보를 비교한다. 생산·기상·달력 추가 실험은 지원 모델끼리 별도 표에 기록한다. 미래 생산량 실적과 미래 실제 기온은 입력할 수 없으며, 예측 당시 확정된 계획·달력·기상예보만 사용한다. |
| 시간순 분할 | 마지막 20%를 최종 시험 구간으로 보관한다. 앞 80%에서 학습 구간이 늘어나는 3개 검증 fold를 만든다. 모든 모델에 같은 날짜 경계를 사용한다. [18, 19] |
| 평가 구간 | 훈련·검증·시험의 정답 시각이 겹치지 않도록 구분한다. 평가 지표는 원래 전력 단위로 계산한다. |
| 피크 정의 | 각 fold 훈련 전력의 95백분위수를 기준값으로 정하고, 실제 관측값이 기준 이상인 평가 시점에서 MAE를 계산한다. 최종 기준값은 시험 이전 훈련 자료로 확정한다. 피크 표본 수를 함께 기록하고, 표본이 없으면 N/A로 표시한다. |
| 설정 선택 | 검증 RMSE를 기본 선택 기준으로 사용한다. 최저 RMSE 대비 2% 이내 후보에서 피크 구간 MAE가 작은 설정을 고르고, 전체 MAE와 시간 비용도 확인한다. 2%는 본 과제의 사전 결정 규칙 제안이며 논문에서 정한 기준은 아니다. |
| 최종 확인 | 모델 종류와 설정은 검증 결과로 선택한다. 선택 완료 후 시험 구간에서 후보별 성능을 한 번 확인하고 결과에 맞춰 탐색 범위를 다시 수정하지 않는다. 시험 결과를 보고 재선정하면 그 구간은 추가 검증 자료가 되므로 새 시험 구간이 필요하다. |

학습 과정의 조기 종료에는 fold 훈련 구간 끝의 별도 시간순 구간을 사용한다. 설정 비교용 검증 구간과 최종 시험 구간을 조기 종료용 자료로 반복 사용하지 않는다. 최종 학습에서는 내부 검증으로 선택한 학습 횟수를 사용해 시험 이전 자료로 다시 학습한다.

### 5.2 XGBoost·LightGBM·LSTM Grid Search

각 표에 제시한 값의 모든 조합을 탐색한다. 공통 입력 길이는 168시간으로 고정해 구조·학습률의 효과를 먼저 비교한다. 입력 길이를 추가 탐색한다면 24·72·168시간을 별도 실험 축으로 두고, 최대 입력 길이에 맞춰 평가 시작 시점을 통일한다.

| 모델 | 탐색 파라미터 | 후보 값 | 확인하려는 효과 |
|---|---|---|---|
| XGBoost | max_depth | 3, 5, 7 | 트리 복잡도와 과적합 |
| XGBoost | learning_rate | 0.01, 0.05, 0.1 | 업데이트 크기 |
| XGBoost | n_estimators | 200, 500, 1000 | 학습률과 트리 수의 조합 |
| XGBoost | min_child_weight | 1, 5 | 작은 표본 집단에 대한 분할 억제 |
| XGBoost | colsample_bytree | 0.8, 1.0 | 입력 변수 일부 사용의 효과 |
| LightGBM | num_leaves | 7, 15, 31 | 잎 개수에 따른 표현력 |
| LightGBM | learning_rate | 0.01, 0.05, 0.1 | 업데이트 크기 |
| LightGBM | n_estimators | 200, 500, 1000 | 학습률과 트리 수의 조합 |
| LightGBM | min_child_samples | 20, 50 | 잎에 필요한 최소 표본 수 |
| LightGBM | colsample_bytree | 0.8, 1.0 | 입력 변수 일부 사용의 효과 |
| LSTM | hidden_size | 32, 64, 128 | 은닉 상태 크기 |
| LSTM | num_layers | 1, 2 | 순환층 깊이 |
| LSTM | learning_rate | 0.0001, 0.0003, 0.001 | 최적화 속도와 안정성 |
| LSTM | head_dropout | 0.0, 0.2 | 마지막 은닉 상태와 출력층 사이의 dropout |

XGBoost는 108개, LightGBM은 108개, LSTM은 36개 조합이다. 3개 fold에서 각각 324회·324회·108회의 설정 평가를 수행한다. 예측 길이 24의 직접 예측 트리 방식에서는 설정·fold마다 24개 회귀기를 학습하므로 트리 회귀기 학습 수는 모델별 7,776회다. 최종 재학습과 반복 seed 확인은 이 수에 추가된다. 데이터 크기만으로 전체 탐색 시간을 단정하지 않고 첫 설정의 실측 시간을 기록한다.

고정 설정은 XGBoost의 objective=reg:squarederror, tree_method=hist, subsample=1.0, reg_lambda=1.0, reg_alpha=0.0이다. LightGBM은 objective=regression, max_depth=-1, subsample=1.0, reg_lambda=1.0, reg_alpha=0.0으로 고정한다. 기본 탐색에서는 n_estimators를 정확한 트리 수로 비교하며 트리 조기 종료를 함께 적용하지 않는다. 깊이·잎 수·최소 표본 수는 공식 문서가 설명하는 복잡도 조절 인자다. [20, 21]

LSTM은 단방향, batch_size=32, Adam, MSE 학습 손실, 최대 100 epoch, 내부 검증 RMSE 기준 patience=10으로 제안한다. num_layers=1에서도 dropout 효과를 비교할 수 있도록 PyTorch LSTM 내부 dropout은 0으로 고정하고 별도 출력부 dropout을 둔다. 내부 dropout은 마지막 층을 제외한 층 사이에 적용되기 때문이다. [22]

XGBoost·LightGBM은 GridSearchCV와 시간순 분할을 사용하고, LSTM은 ParameterGrid로 동일한 전체 조합을 순회하는 학습 루프를 사용한다. 탐색 시 seed=42를 고정하고, 모델별 상위 설정 3개는 seed=17·42·73으로 확인한다. 평균과 표준편차는 fold 간 변동과 seed 간 변동을 구분해 기록한다. 그리드 가장자리 값이 가장 좋으면 시험 구간을 열기 전에 해당 축을 확장한다.

### 5.3 사전학습 모델의 층별 구성 확인

층 비교의 단위는 attention·feed-forward·정규화를 포함하는 Transformer 블록이다. 블록 내부의 개별 선형층 개수로 세지 않는다. 아래 경로는 핵심 모델 객체 기준이며 외부 pipeline의 접두 경로는 제외했다.

| 모델 | 확인된 블록 수 / 경로 | 출력부 | 미세조정 구현 상태 |
|---|---|---|---|
| Chronos-2 | 12개 / encoder.block | output_patch_embedding + encoder.final_layer_norm | 공식 fit은 full·LoRA를 지원한다. 층별 동결 옵션은 별도 구현이 필요하다. [23–25] |
| TimesFM 3.0 | 20개 / transformer_stack.layers | output_head | 공식 PyTorch 코드는 inference only로 명시됐다. forward는 공개되어 있으나 학습 손실·학습 루프의 별도 구현과 검증이 필요하다. [26–28] |
| Moirai 2.0-R-small | 6개 / encoder.layers | out_proj + encoder.norm | 모듈 forward에 training_mode가 있으나 확인된 일반 fine-tuning 예제는 Moirai 1 계열이다. 2.0용 손실·학습 래퍼를 연결해야 한다. [29–32] |

Chronos-2의 fit은 모델을 새로 생성하고 가중치를 복사한다. 호출 전 원본 모델의 requires_grad만 바꾸면 동결 설정이 새 모델에 이어지지 않을 수 있다. 복사한 실제 학습 모델에서 층을 지정하고 optimizer를 만들기 전에 동결을 적용해야 한다. [25]

TimesFM 3.0의 decode는 no_grad 경로이므로 이를 그대로 학습 루프로 사용하지 않는다. forward를 이용하는 학습 경로에서 역전파를 검증해야 한다. 기존 2.5용 LoRA 예제를 3.0에 그대로 적용했다고 기록할 수 없다. [27]

Moirai 2.0은 전력 이력 공통 비교에 사용한다. 생산 변수를 공동 입력하는 확장 비교에는 포함하지 않는다. [10, 30]

### 5.4 모델별 9개 학습 범위

각 실험은 동일한 원본 사전학습 가중치에서 독립적으로 시작한다. 한 실험의 학습 결과를 다음 실험의 초기값으로 넘기지 않는다. ZS는 가중치를 업데이트하지 않으며, F0–F7은 학습하는 파라미터 범위만 다르게 설정한다. 출력부에는 위 표의 마지막 정규화층도 포함한다.

| 설정 | Chronos-2 (12블록) | TimesFM 3.0 (20블록) | Moirai 2.0 (6블록) |
|---|---|---|---|
| ZS | Zero-shot | Zero-shot | Zero-shot |
| F0 | 출력부만 | 출력부만 | 출력부만 |
| F1 | 마지막 1블록 + 출력부 | 마지막 1블록 + 출력부 | 마지막 1블록 + 출력부 |
| F2 | 마지막 2블록 + 출력부 | 마지막 2블록 + 출력부 | 마지막 2블록 + 출력부 |
| F3 | 마지막 3블록 + 출력부 | 마지막 5블록 + 출력부 | 마지막 3블록 + 출력부 |
| F4 | 마지막 6블록 + 출력부 | 마지막 10블록 + 출력부 | 마지막 4블록 + 출력부 |
| F5 | 마지막 9블록 + 출력부 | 마지막 15블록 + 출력부 | 마지막 5블록 + 출력부 |
| F6 | 12블록 + 출력부 | 20블록 + 출력부 | 6블록 + 출력부 |
| F7 | 입력부를 포함한 전체 | 입력부를 포함한 전체 | 입력부를 포함한 전체 |

F0–F6에서는 입력 임베딩·입력 투영부를 고정한다. F6과 F7은 전체 블록을 학습한다는 점은 같지만 F7에서만 입력부 등 남은 파라미터도 학습한다. 이 구분으로 입력 변환까지 바꿀 필요가 있는지 확인한다. 마지막 k개 블록은 0부터 시작하는 인덱스에서 [전체 블록 수−k, 전체 블록 수−1]이다. F0의 k=0은 별도 처리해 전체 블록이 선택되지 않게 한다.

학습률은 각 F 설정에서 1e−6·1e−5·1e−4의 동일한 세 값을 탐색한다. context=168, horizon=24, 유효 batch_size=32, 최대 1,000 optimizer step, 100 step마다 내부 검증, 개선 없는 3회 검증 후 종료를 초기 조건으로 제안한다. 검증 RMSE가 계속 감소하면 최적점이 관측되지 않은 것으로 보고 시험 평가 전에 모든 범위에 같은 확대 규칙을 적용한다. 모델마다 실제 메모리에 맞춰 작은 batch와 gradient accumulation을 사용하고 유효 batch는 유지한다.

층 수만의 영향을 보는 그래프는 같은 학습률의 곡선을 함께 남긴다. 범위별 최고 결과를 보여주는 곡선에는 각 범위에서 검증으로 선택한 학습률을 표기한다. 두 종류를 구분해야 학습률 효과를 층 효과로 오해하지 않는다. 각 모델의 학습 손실은 범위 간 동일하게 유지한다. Chronos-2는 공식 학습 손실을 유지하고, 별도 학습 구현이 필요한 두 모델은 중앙값 출력의 MSE를 사용하는 점예측 적응안을 제안한다. 이 손실 선택은 원 논문 재현 조건이 아니라 본 과제의 제안 조건으로 기록한다.

미세조정이 구현된 모델마다 8범위 × 3학습률 × 3fold = 72회 학습을 수행한다. ZS는 별도로 같은 평가 구간에서 추론한다. 검증 상위 범위·학습률 3개는 seed=17·42·73으로 재확인한다. 모델별 층 수 비율은 파라미터 수 비율과 다르므로 실제 학습 파라미터 개수·비율도 기록한다. LoRA는 이번 범위 비교에 섞지 않고 필요할 때 독립된 추가 실험으로 둔다.

실행 전 한 batch에서 유한한 손실과 gradient가 나오는지, 지정한 파라미터만 업데이트되는지, 동결 파라미터가 그대로인지, 저장 후 예측이 재현되는지 확인한다. zero-shot 학습 래퍼와 공식 추론의 출력 정렬도 확인한다. 이 확인을 통과한 구현만 층별 결과표에 포함한다. TimesFM 3.0·Moirai 2.0의 층별 설정은 해당 학습 구현 검증을 전제로 한 실험안이다.

## 6. 결과 그래프 예시와 모델 선정에 사용할 산출물

이 절의 모든 그래프는 결과물의 형태와 해석 방법을 보여주는 가상 데이터다. 실제 학습 결과, 문헌 성능 수치, 후보별 예상 순위를 나타내지 않는다. 실제 값은 실험 후 원래 전력 단위로 대체한다. 특정 모델이 반드시 우세하거나 전체 미세조정이 항상 좋아진다고 가정하지 않는다.

### 6.1 Grid Search 결과

![그림 1. 가상 데이터로 만든 Grid Search 결과 예시](report_assets/01_grid_search_example.png)

각 칸은 특정 파라미터 조합의 3개 검증 fold 평균 RMSE다. 그림은 나머지 파라미터를 고정한 단면이다. 실제 탐색에서는 모든 조합을 CSV로 보관하고, 선택한 단면의 고정값도 제목에 기록한다. 최저 오차 주변에도 좋은 조합이 모이는지, 최적값이 탐색 범위 가장자리에 있는지 확인한다. 표준편차가 큰 조합은 추가 seed 결과와 함께 판단한다.

### 6.2 학습 층 범위에 따른 추세

![그림 2. 가상 데이터로 만든 층별 학습 범위 비교](report_assets/02_layer_trends_example.png)

세 패널은 각각 과적합 가능성, 계속되는 개선, 성능 정체라는 서로 다른 가상 추세를 보여준다. 실제 모델의 반응을 예측한 그림은 아니다. 훈련 오차만 줄고 검증 오차가 커지면 더 많은 층을 학습하는 설정을 선택할 근거가 약하다. 검증 오차가 비슷하면 학습 시간이 짧고 변동이 작은 범위를 고려한다. ZS의 훈련 오차는 정의하지 않으며, 오차막대는 표준편차로 표시하고 신뢰구간이라고 부르지 않는다.

### 6.3 모델별 전체·피크 오차 비교

![그림 3. 가상 데이터로 만든 모델별 오차 비교](report_assets/03_model_comparison_example.png)

각 모델의 검증으로 선택한 설정을 한 행에 두고, 사전학습 모델은 ZS와 선택된 FT를 각각 기록한다. 전체 RMSE 순위와 피크 MAE 순위가 다를 수 있으므로 세 지표를 함께 본다. 공통 전력 이력 실험과 외생 변수 확장 실험은 그림을 분리한다. 구현 검증을 통과하지 못한 FT는 실제 결과표에 수치를 채우지 않는다. 최종 시험 그래프는 선정 완료 후 확인용으로 작성한다.

### 6.4 실제값과 예측값의 시간별 비교

![그림 4. 합성 시계열의 피크 예측 비교](report_assets/04_peak_trace_example.png)

위 패널에서 피크 높이와 발생 시점이 맞는지 확인하고, 아래 패널에서 음수 오차가 피크 주변에 집중되는지 확인한다. 음수는 과소예측이다. 실제 그래프는 전체 시험 구간과 사전 정의한 피크 사건 창을 모두 제공한다. 예측이 잘 맞는 구간만 골라 최종 선정을 설명하지 않는다. 이 예시의 120 기준은 가상 값이며 실제 기준은 훈련 자료에서 계산한다.

### 6.5 정확도와 학습 비용 비교

![그림 5. 가상 데이터로 만든 학습 시간과 오차 비교](report_assets/05_cost_accuracy_example.png)

이 그림은 선택된 설정 1회의 추가 학습 시간과 검증 RMSE를 비교한다. Zero-shot은 추가 학습 시간이 0이지만 로딩·추론 시간은 발생한다. 실제 결과에는 전체 탐색 시간, 선택 설정 재학습 시간, 동일 장비의 추론 지연을 별도로 기록한다. 두 모델이 모두 더 낮은 오차와 더 짧은 시간을 갖는지 확인하고, 성능 차이가 작으면 운영 조건에 맞는 비용을 검토한다.

### 6.6 모델 선정용 결과표

| 산출물 | 반드시 기록할 내용 | 선정에 사용하는 방법 |
|---|---|---|
| 모든 탐색 결과 CSV | 모델·입력 조건·fold·seed·파라미터·학습 범위·학습률·검증 MAE/RMSE/피크 MAE·피크 표본 수 | 최저값 하나뿐 아니라 안정적인 설정인지 확인 |
| 예측값 CSV | 예측 시작 시각·정답 시각·예측 시차·관측값·예측값·피크 여부·모델·실험 ID | 같은 대상 시각을 예측했는지 확인하고 오차 재계산 |
| 학습 범위 기록 | 학습한 모듈 이름·블록 인덱스·학습 가능/전체 파라미터 수·학습률·step·손실 | 층별 변화와 학습률 변화의 원인 구분 |
| 후보별 요약표 | 검증 세 지표·fold/seed별 변동·시험 세 지표·ZS 대비 FT 변화 | 검증으로 선정하고 최종 시험에서 일반화 성능 확인 |
| 비용 및 재현 기록 | 전체 탐색/선택 설정 학습/추론 시간·장비·최대 메모리·라이브러리 버전·가중치 revision | 반복 실행과 사용 환경의 실행 가능성 확인 |

선정 순서는 ① 공통 입력·동일 시각 평가 확인 → ② 검증 RMSE 최저 대비 2% 이내 후보 확인 → ③ 피크 MAE와 전체 MAE·변동 비교 → ④ 학습·추론 비용 및 기존 라이선스 조건 확인 → ⑤ 모델·설정 확정 → ⑥ 최종 시험 성능 보고다. 그래프의 가상 수치로 우선 후보를 정하지 않는다.

## 7. 실험 설계 추가 근거

[18] scikit-learn, GridSearchCV. https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.GridSearchCV.html

[19] scikit-learn, TimeSeriesSplit. https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.TimeSeriesSplit.html

[20] XGBoost, Parameters. https://xgboost.readthedocs.io/en/stable/parameter.html

[21] LightGBM, Parameters Tuning. https://lightgbm.readthedocs.io/en/stable/Parameters-Tuning.html

[22] PyTorch, LSTM. https://docs.pytorch.org/docs/2.9/generated/torch.nn.LSTM.html

[23] Amazon, Chronos-2 checkpoint configuration. https://huggingface.co/amazon/chronos-2/blob/main/config.json

[24] Amazon Science, Chronos2Model / Chronos2Encoder. https://github.com/amazon-science/chronos-forecasting/blob/main/src/chronos/chronos2/model.py

[25] Amazon Science, Chronos2Pipeline.fit. https://github.com/amazon-science/chronos-forecasting/blob/main/src/chronos/chronos2/pipeline.py

[26] Google Research, TimesFM 3.0 checkpoint configuration. https://huggingface.co/google/timesfm-3.0-pytorch/blob/main/config.json

[27] Google Research, TimesFM3Torch. https://github.com/google-research/timesfm/blob/master/src/timesfm3/torch/model.py

[28] Google Research, StackedMixingTransformer. https://github.com/google-research/timesfm/blob/master/src/timesfm3/torch/transformer.py

[29] Salesforce, Moirai 2.0-R-small checkpoint configuration. https://huggingface.co/Salesforce/moirai-2.0-R-small/blob/main/config.json

[30] Salesforce AI Research, Moirai2Module. https://github.com/SalesforceAIResearch/uni2ts/blob/main/src/uni2ts/model/moirai2/module.py

[31] Salesforce AI Research, TransformerEncoder. https://github.com/SalesforceAIResearch/uni2ts/blob/main/src/uni2ts/module/transformer.py

[32] Salesforce AI Research, Uni2TS fine-tuning example. https://github.com/SalesforceAIResearch/uni2ts/blob/main/README.md#fine-tuning
'''


def plain(text):
    return text.replace('**', '').replace('`', '')


def add_md(doc, text):
    lines = text.strip().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('|'):
            data = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [plain(x.strip()) for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', x.replace(' ', '')) for x in row): data.append(row)
                i += 1
            table = doc.add_table(rows=1, cols=len(data[0]))
            table.style = 'Table Grid'
            for j,x in enumerate(data[0]): table.rows[0].cells[j].text = x
            for row in data[1:]:
                cells = table.add_row().cells
                for j,x in enumerate(row): cells[j].text = x
            # Keep rows intact, repeat the header, and use consistent compact fonts.
            for n,row in enumerate(table.rows):
                tr_pr = row._tr.get_or_add_trPr()
                cant = OxmlElement('w:cantSplit'); tr_pr.append(cant)
                if n == 0:
                    repeat = OxmlElement('w:tblHeader'); tr_pr.append(repeat)
                for cell in row.cells:
                    if n == 0:
                        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EFF7')
                        cell._tc.get_or_add_tcPr().append(shade)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(3)
                        p.paragraph_format.keep_with_next = n == 0
                        for run in p.runs:
                            run.font.size = Pt(8)
                            run.bold = n == 0
            doc.add_paragraph()
            continue
        if line.startswith('!['):
            match = re.match(r'!\[(.*?)\]\((.*?)\)',line)
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(ROOT / match[2]), width=Inches(6.9))
            cap = doc.add_paragraph(match[1] + ' — ' + NOTE)
            cap.paragraph_format.space_after = Pt(6)
            for run in cap.runs:
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor.from_string('A05A13')
        elif line.startswith('### '):
            p = doc.add_heading(line[4:], level=2)
            if re.match(r'6\.[2-5] ', line[4:]): p.paragraph_format.page_break_before = True
        elif line.startswith('## '):
            p = doc.add_heading(line[3:],level=1)
            p.paragraph_format.page_break_before = True
        else:
            p = doc.add_paragraph(plain(line))
            p.paragraph_format.space_after = Pt(6)
        i += 1


def update_documents():
    readme = ROOT / 'README.md'
    original_md = readme.read_text(encoding='utf-8').split(MARKER)[0].rstrip()
    new_md = original_md + '\n\n' + APPENDIX.strip() + '\n'
    readme.write_text(new_md, encoding='utf-8')
    path = ROOT / f'{TITLE}.docx'
    doc = Document(path)
    # Idempotent append: preserve all existing report content preceding section 5.
    remove = False
    for element in list(doc._element.body):
        if element.tag == qn('w:sectPr'): continue
        if element.tag == qn('w:p'):
            content = ''.join(element.itertext())
            if '5. 하이퍼파라미터 탐색 및 층별 미세조정 실험 설계' in content: remove = True
        if remove: doc._element.body.remove(element)
    add_md(doc, APPENDIX)
    doc.save(path)
    docs = ROOT / 'docs'
    if docs.exists():
        (docs / f'{TITLE}.md').write_text(new_md.replace('](report_assets/', '](../report_assets/'),encoding='utf-8')
        shutil.copy2(path, docs / path.name)


if __name__ == '__main__':
    enumerate_plan()
    make_figures()
    update_documents()
    print('Updated README and DOCX; created 5 synthetic figures and planned-trial files. No model training performed.')
